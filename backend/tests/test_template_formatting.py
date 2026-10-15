
from docx import Document
from docx.oxml.ns import qn

from documents.templates_engine.spacing import apply_spacing


def test_spacing_applies_to_body_paragraphs():
    doc = Document()
    first = doc.add_paragraph("Primer párrafo")
    second = doc.add_paragraph("Segundo párrafo")

    first.paragraph_format.line_spacing = 2.0

    apply_spacing(doc, 1.5)

    assert doc.styles["Normal"].paragraph_format.line_spacing == 1.5

    for paragraph in (first, second):
        paragraph_format = paragraph.paragraph_format

        assert paragraph_format.line_spacing == 1.5

        spacing_element = (
            paragraph_format._element
            .get_or_add_pPr()
            .get_or_add_spacing()
        )

        # "auto" significa interlineado múltiple en OOXML.
        assert spacing_element.get(qn("w:lineRule")) == "auto"

        # 360 equivale a 1.5 líneas.
        assert spacing_element.get(qn("w:line")) == "360"

    normal_spacing = (
        doc.styles["Normal"].paragraph_format._element
        .get_or_add_pPr()
        .get_or_add_spacing()
    )

    assert normal_spacing.get(qn("w:lineRule")) == "auto"
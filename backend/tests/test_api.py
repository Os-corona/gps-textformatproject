import io
import os
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document as DocxDocument
from rest_framework import status
from rest_framework.test import APITestCase

from documents.models import Document, ProcessingJob


FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"
INVALID_FIXTURES_DIR = FIXTURES_DIR / "invalid"


class DocumentAPITests(APITestCase):
    def setUp(self):
        self.valid_docx_path = FIXTURES_DIR / "01_planos.docx"

    def _get_valid_uploaded_file(self):
        with open(self.valid_docx_path, "rb") as f:
            return SimpleUploadedFile(
                "01_planos.docx",
                f.read(),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

    # -------------------------------------------------------------------------
    # 2) Test de carga válida: responde 201 con los ids.
    # -------------------------------------------------------------------------
    def test_valid_upload_returns_201_with_ids(self):
        """
        Valida que subir un archivo .docx válido responda 201 Created
        e incluya document_id, job_id y status 'pending'.
        """
        uploaded_file = self._get_valid_uploaded_file()

        response = self.client.post(
            "/api/documents/",
            {"file": uploaded_file},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIn("document_id", response.data)
        self.assertIn("job_id", response.data)
        self.assertIsInstance(response.data["document_id"], int)
        self.assertIsInstance(response.data["job_id"], int)
        self.assertEqual(response.data["status"], "pending")
        self.assertEqual(response.data["original_name"], "01_planos.docx")

        # Verificar persistencia en base de datos
        self.assertTrue(Document.objects.filter(id=response.data["document_id"]).exists())
        self.assertTrue(ProcessingJob.objects.filter(id=response.data["job_id"]).exists())

    def test_valid_upload_via_upload_url(self):
        """
        Verifica compatibilidad con la ruta alternativa /api/documents/upload/.
        """
        uploaded_file = self._get_valid_uploaded_file()

        response = self.client.post(
            "/api/documents/upload/",
            {"file": uploaded_file},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIn("document_id", response.data)
        self.assertIn("job_id", response.data)

    # -------------------------------------------------------------------------
    # 3) Tests de carga inválida: falso.docx, vacio.docx, corrupto.docx,
    #    un archivo mayor a 10 MB y un .pdf. Todos responden 400 con el código esperado.
    # -------------------------------------------------------------------------
    def test_invalid_upload_falso_docx(self):
        """
        Archivo con extensión .docx pero con contenido PDF (falso.docx).
        Debe responder 400 con error INVALID_FILE.
        """
        falso_path = INVALID_FIXTURES_DIR / "falso.docx"
        with open(falso_path, "rb") as f:
            uploaded_file = SimpleUploadedFile("falso.docx", f.read())

        response = self.client.post(
            "/api/documents/",
            {"file": uploaded_file},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("error"), "INVALID_FILE")

    def test_invalid_upload_vacio_docx(self):
        """
        Archivo ZIP vacío de 22 bytes sin estructura OOXML (vacio.docx).
        Debe responder 400 con error INVALID_FILE.
        """
        vacio_path = INVALID_FIXTURES_DIR / "vacio.docx"
        with open(vacio_path, "rb") as f:
            uploaded_file = SimpleUploadedFile("vacio.docx", f.read())

        response = self.client.post(
            "/api/documents/",
            {"file": uploaded_file},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("error"), "INVALID_FILE")

    def test_invalid_upload_corrupto_docx(self):
        """
        Archivo DOCX truncado/corrupto (corrupto.docx).
        Debe responder 400 con error INVALID_FILE.
        """
        corrupto_path = INVALID_FIXTURES_DIR / "corrupto.docx"
        with open(corrupto_path, "rb") as f:
            uploaded_file = SimpleUploadedFile("corrupto.docx", f.read())

        response = self.client.post(
            "/api/documents/",
            {"file": uploaded_file},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("error"), "INVALID_FILE")

    def test_invalid_upload_file_larger_than_10mb(self):
        """
        Archivo mayor a 10 MB (10 * 1024 * 1024 + 1024 bytes).
        Debe responder 400 con error FILE_TOO_LARGE.
        """
        large_size = 10 * 1024 * 1024 + 1024
        # Generar contenido que exceda los 10 MB con encabezado ZIP
        large_content = b"PK\x03\x04" + b"0" * (large_size - 4)
        uploaded_file = SimpleUploadedFile(
            "archivo_grande.docx",
            large_content,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        response = self.client.post(
            "/api/documents/",
            {"file": uploaded_file},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("error"), "FILE_TOO_LARGE")

    def test_invalid_upload_pdf(self):
        """
        Archivo con extensión .pdf.
        Debe responder 400 con error INVALID_FILE.
        """
        uploaded_file = SimpleUploadedFile(
            "documento.pdf",
            b"%PDF-1.4 contenido de prueba",
            content_type="application/pdf",
        )

        response = self.client.post(
            "/api/documents/",
            {"file": uploaded_file},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("error"), "INVALID_FILE")

    # -------------------------------------------------------------------------
    # 4) Tests de configuración fuera de rango (margen de 20 cm, tamaño de 2 pt,
    #    fuente no permitida): responde 400 INVALID_CONFIG.
    # -------------------------------------------------------------------------
    def test_format_out_of_range_margin(self):
        """
        Configuración con margen de 20 cm (el límite permitido es 0.5 a 5.0 cm).
        Debe responder 400 INVALID_CONFIG.
        """
        upload_resp = self.client.post(
            "/api/documents/",
            {"file": self._get_valid_uploaded_file()},
            format="multipart",
        )
        doc_id = upload_resp.data["document_id"]

        invalid_config = {
            "preset_id": "apa7",
            "margins_cm": {"top": 20.0},
        }

        response = self.client.post(
            f"/api/documents/{doc_id}/format/",
            invalid_config,
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("error"), "INVALID_CONFIG")

    def test_format_out_of_range_font_size(self):
        """
        Configuración con tamaño de fuente de 2 pt (el límite permitido es 8 a 24 pt).
        Debe responder 400 INVALID_CONFIG.
        """
        upload_resp = self.client.post(
            "/api/documents/",
            {"file": self._get_valid_uploaded_file()},
            format="multipart",
        )
        doc_id = upload_resp.data["document_id"]

        invalid_config = {
            "preset_id": "apa7",
            "font_size_pt": 2,
        }

        response = self.client.post(
            f"/api/documents/{doc_id}/format/",
            invalid_config,
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("error"), "INVALID_CONFIG")

    def test_format_disallowed_font(self):
        """
        Configuración con una fuente no permitida (e.g. Comic Sans MS).
        Debe responder 400 INVALID_CONFIG.
        """
        upload_resp = self.client.post(
            "/api/documents/",
            {"file": self._get_valid_uploaded_file()},
            format="multipart",
        )
        doc_id = upload_resp.data["document_id"]

        invalid_config = {
            "preset_id": "apa7",
            "font_family": "Comic Sans MS",
        }

        response = self.client.post(
            f"/api/documents/{doc_id}/format/",
            invalid_config,
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("error"), "INVALID_CONFIG")

    def test_format_nonexistent_document_returns_404(self):
        """
        Intentar formatear un document_id inexistente debe responder 404 NOT_FOUND.
        """
        response = self.client.post(
            "/api/documents/999999/format/",
            {"preset_id": "apa7"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(response.data.get("error"), "NOT_FOUND")

    # -------------------------------------------------------------------------
    # 5) Tests de descarga: id inexistente responde 404 y descarga antes de
    #    formatear responde 409.
    # -------------------------------------------------------------------------
    def test_download_nonexistent_id_returns_404(self):
        """
        Intentar descargar un documento inexistente debe responder 404 NOT_FOUND.
        """
        response = self.client.get("/api/documents/999999/download/")

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(response.data.get("error"), "NOT_FOUND")

    def test_download_before_formatting_returns_409(self):
        """
        Intentar descargar un documento que aún no ha sido formateado (status='pending')
        debe responder 409 NOT_READY.
        """
        upload_resp = self.client.post(
            "/api/documents/",
            {"file": self._get_valid_uploaded_file()},
            format="multipart",
        )
        doc_id = upload_resp.data["document_id"]

        response = self.client.get(f"/api/documents/{doc_id}/download/")

        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(response.data.get("error"), "NOT_READY")

    # -------------------------------------------------------------------------
    # 6) Test de flujo completo: carga → formateo → consulta de estado done →
    #    descarga → abrir los bytes con python-docx y verificar margen y fuente.
    # -------------------------------------------------------------------------
    def test_complete_flow_upload_format_status_download_and_verify(self):
        """
        Flujo completo:
        1. Carga del documento (POST /api/documents/) -> 201 Created
        2. Formateo (POST /api/documents/{id}/format/) con márgenes y fuente específicos -> 200 OK
        3. Consulta de estado (GET /api/documents/{id}/) -> 200 OK y status 'done'
        4. Descarga (GET /api/documents/{id}/download/) -> 200 OK archivo binario DOCX
        5. Abrir bytes descargados con python-docx y verificar márgenes y fuente aplicada
        """
        # 1. Carga
        uploaded_file = self._get_valid_uploaded_file()
        upload_resp = self.client.post(
            "/api/documents/",
            {"file": uploaded_file},
            format="multipart",
        )
        self.assertEqual(upload_resp.status_code, status.HTTP_201_CREATED)
        doc_id = upload_resp.data["document_id"]
        job_id = upload_resp.data["job_id"]
        self.assertIsNotNone(doc_id)
        self.assertIsNotNone(job_id)

        # 2. Formateo con márgenes de 3.0 cm y fuente Arial de 12 pt
        target_margin = 3.0
        target_font = "Arial"
        format_config = {
            "preset_id": "apa7",
            "margins_cm": {
                "top": target_margin,
                "bottom": target_margin,
                "left": target_margin,
                "right": target_margin,
            },
            "font_family": target_font,
            "font_size_pt": 12,
            "line_spacing": 1.5,
        }
        format_resp = self.client.post(
            f"/api/documents/{doc_id}/format/",
            format_config,
            format="json",
        )
        self.assertEqual(format_resp.status_code, status.HTTP_200_OK)
        self.assertEqual(format_resp.data.get("status"), "done")

        # 3. Consulta de estado
        status_resp = self.client.get(f"/api/documents/{doc_id}/")
        self.assertEqual(status_resp.status_code, status.HTTP_200_OK)
        self.assertEqual(status_resp.data.get("status"), "done")
        self.assertEqual(status_resp.data.get("document_id"), doc_id)
        self.assertEqual(
            status_resp.data.get("download_url"),
            f"/api/documents/{doc_id}/download/",
        )

        # 4. Descarga del documento resultante
        download_resp = self.client.get(f"/api/documents/{doc_id}/download/")
        self.assertEqual(download_resp.status_code, status.HTTP_200_OK)
        self.assertEqual(
            download_resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("01_planos_formateado.docx", download_resp["Content-Disposition"])

        if hasattr(download_resp, "streaming_content"):
            download_bytes = b"".join(download_resp.streaming_content)
        else:
            download_bytes = download_resp.content

        self.assertGreater(len(download_bytes), 0)

        # 5. Abrir los bytes con python-docx
        result_doc = DocxDocument(io.BytesIO(download_bytes))
        self.assertIsNotNone(result_doc)

        # 6. Verificar margen y fuente
        # Verificar márgenes en todas las secciones
        self.assertGreater(len(result_doc.sections), 0)
        for i, section in enumerate(result_doc.sections):
            self.assertEqual(
                round(section.top_margin.cm, 2),
                target_margin,
                f"Sección {i} top_margin incorrecto",
            )
            self.assertEqual(
                round(section.bottom_margin.cm, 2),
                target_margin,
                f"Sección {i} bottom_margin incorrecto",
            )
            self.assertEqual(
                round(section.left_margin.cm, 2),
                target_margin,
                f"Sección {i} left_margin incorrecto",
            )
            self.assertEqual(
                round(section.right_margin.cm, 2),
                target_margin,
                f"Sección {i} right_margin incorrecto",
            )

        # Verificar fuente en el estilo Normal y en los runs
        self.assertEqual(result_doc.styles["Normal"].font.name, target_font)
        self.assertGreater(len(result_doc.paragraphs), 0)
        for p in result_doc.paragraphs:
            for run in p.runs:
                if run.text.strip():
                    self.assertEqual(
                        run.font.name,
                        target_font,
                        f"Run con texto '{run.text}' no tiene la fuente configurada",
                    )

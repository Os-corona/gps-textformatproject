"""
Pruebas unitarias y de integración para templates_engine (S2T15).

Valida la aplicación de plantillas de formato (márgenes, tipografía, interlineado)
y la preservación de elementos estructurales (tablas, imágenes, listas) sobre
la suite de documentos DOCX de prueba (fixtures).
"""

import io
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from documents.templates_engine.applier import apply_template
from documents.templates_engine.config import TemplateConfig

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"

# Cargar únicamente los documentos DOCX válidos directamente dentro de fixtures/
# (excluyendo la subcarpeta invalid/ y archivos temporales de Word con prefijo ~$)
VALID_FIXTURE_PATHS = sorted([
    p for p in FIXTURES_DIR.glob("*.docx")
    if p.is_file() and not p.name.startswith("~$")
])


@pytest.fixture(params=VALID_FIXTURE_PATHS, ids=lambda p: p.name)
def valid_docx_path(request):
    """Fixture parametrizada que provee la ruta a cada archivo .docx válido de fixtures/."""
    return request.param


@pytest.fixture
def loaded_doc(valid_docx_path):
    """Provee una instancia fresca de Document para cada prueba."""
    return Document(str(valid_docx_path))


@pytest.fixture
def sample_config():
    """Configuración estándar de prueba basada en formato APA 7."""
    return TemplateConfig(
        margin_top=2.54,
        margin_bottom=2.54,
        margin_left=2.54,
        margin_right=2.54,
        font_name="Times New Roman",
        font_size=12,
        line_spacing=2.0,
    )


def count_list_paragraphs(doc: Document) -> int:
    """
    Cuenta los párrafos que constituyen elementos de lista.
    Detecta listas estándar de Word (w:numPr), estilos con nombre de lista ('List')
    y sangría francesa con viñeta (w:ind[@w:hanging]).
    """
    count = 0
    for p in doc.paragraphs:
        if p._p.xpath("./w:pPr/w:numPr"):
            count += 1
        elif p.style and "list" in p.style.name.lower():
            count += 1
        elif p._p.xpath("./w:pPr/w:ind[@w:hanging]"):
            count += 1
    return count


def iter_all_runs(doc: Document):
    """
    Generador que recorre todos los runs del cuerpo del documento
    y de las celdas de las tablas.
    """
    for p in doc.paragraphs:
        for r in p.runs:
            yield ("párrafo del cuerpo", p, r)

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        yield ("celda de tabla", p, r)


# ---------------------------------------------------------------------------
# Requisito 3: Test de márgenes
# ---------------------------------------------------------------------------
def test_margins_applied_to_all_sections(loaded_doc, sample_config):
    """
    Verifica que tras apply_template, todas las secciones del documento
    tengan configurados los márgenes especificados (superior, inferior, izquierdo, derecho).
    """
    apply_template(loaded_doc, sample_config)

    assert len(loaded_doc.sections) > 0, "El documento debe contener al menos una sección."

    for i, section in enumerate(loaded_doc.sections):
        assert round(section.top_margin.cm, 2) == sample_config.margin_top, (
            f"Sección {i}: margen superior esperado {sample_config.margin_top} cm, "
            f"obtenido {round(section.top_margin.cm, 2)} cm"
        )
        assert round(section.bottom_margin.cm, 2) == sample_config.margin_bottom, (
            f"Sección {i}: margen inferior esperado {sample_config.margin_bottom} cm, "
            f"obtenido {round(section.bottom_margin.cm, 2)} cm"
        )
        assert round(section.left_margin.cm, 2) == sample_config.margin_left, (
            f"Sección {i}: margen izquierdo esperado {sample_config.margin_left} cm, "
            f"obtenido {round(section.left_margin.cm, 2)} cm"
        )
        assert round(section.right_margin.cm, 2) == sample_config.margin_right, (
            f"Sección {i}: margen derecho esperado {sample_config.margin_right} cm, "
            f"obtenido {round(section.right_margin.cm, 2)} cm"
        )


# ---------------------------------------------------------------------------
# Requisito 4: Test de tipografía (Estilos y Runs)
# ---------------------------------------------------------------------------
def test_typography_styles_no_theme_attributes(loaded_doc, sample_config):
    """
    Comprueba que los estilos Normal y Heading 1 a 3 usen la fuente configurada
    sin atributos de tema en w:rFonts (e.g. w:asciiTheme, w:hAnsiTheme).
    """
    apply_template(loaded_doc, sample_config)

    target_styles = ["Normal", "Heading 1", "Heading 2", "Heading 3"]
    for style_name in target_styles:
        if style_name in loaded_doc.styles:
            style = loaded_doc.styles[style_name]
            # Verificar nombre de fuente asignado al estilo
            assert style.font.name == sample_config.font_name, (
                f"Estilo '{style_name}' tiene fuente '{style.font.name}', "
                f"se esperaba '{sample_config.font_name}'."
            )

            # Verificar ausencia de atributos de tema en el XML w:rFonts
            rPr = style.element.find(qn("w:rPr"))
            if rPr is not None:
                rFonts = rPr.find(qn("w:rFonts"))
                if rFonts is not None:
                    theme_attributes = [
                        attr for attr in rFonts.attrib
                        if "theme" in attr.lower()
                    ]
                    assert not theme_attributes, (
                        f"Estilo '{style_name}' contiene atributos de tema: {theme_attributes}"
                    )


def test_typography_runs_no_divergent_font_or_size(loaded_doc, sample_config):
    """
    Recorre todos los runs (cuerpo y celdas de tablas) comprobando que ninguno
    tenga fuente o tamaño propio distinto al configurado.
    """
    apply_template(loaded_doc, sample_config)

    for origin, paragraph, run in iter_all_runs(loaded_doc):
        # Si el run tiene fuente explícita directa, debe coincidir con la configurada
        if run.font.name is not None:
            assert run.font.name == sample_config.font_name, (
                f"Run en {origin} ('{run.text[:20]}') tiene fuente '{run.font.name}', "
                f"distinta a la configurada '{sample_config.font_name}'."
            )

        # Si el run contiene elemento w:rFonts, no debe tener atributos de tema
        if run._r.rPr is not None:
            rFonts = run._r.rPr.find(qn("w:rFonts"))
            if rFonts is not None:
                theme_attributes = [
                    attr for attr in rFonts.attrib
                    if "theme" in attr.lower()
                ]
                assert not theme_attributes, (
                    f"Run en {origin} ('{run.text[:20]}') tiene atributos de tema en w:rFonts: {theme_attributes}"
                )

        # Si el run tiene tamaño explícito, debe coincidir con el configurado
        if run.font.size is not None:
            expected_size = sample_config.font_size
            style_name = paragraph.style.name if paragraph.style else ""
            heading_sizes = getattr(sample_config, "heading_sizes", None) or {}

            if "Heading 1" in style_name and "h1" in heading_sizes:
                expected_size = heading_sizes["h1"]
            elif "Heading 2" in style_name and "h2" in heading_sizes:
                expected_size = heading_sizes["h2"]
            elif "Heading 3" in style_name and "h3" in heading_sizes:
                expected_size = heading_sizes["h3"]

            assert round(run.font.size.pt, 1) == expected_size, (
                f"Run en {origin} ('{run.text[:20]}') tiene tamaño {round(run.font.size.pt, 1)} pt, "
                f"distinto al configurado {expected_size} pt."
            )


# ---------------------------------------------------------------------------
# Requisito 5: Test de interlineado en los párrafos del cuerpo
# ---------------------------------------------------------------------------
def test_line_spacing_applied_to_body_paragraphs(loaded_doc, sample_config):
    """
    Verifica que el interlineado múltiple configurado se aplique
    a todos los párrafos del cuerpo del documento.
    """
    apply_template(loaded_doc, sample_config)

    for i, paragraph in enumerate(loaded_doc.paragraphs):
        # 1. Comprobación a nivel de API de python-docx
        assert paragraph.paragraph_format.line_spacing == pytest.approx(sample_config.line_spacing), (
            f"Párrafo {i} ('{paragraph.text[:25]}') tiene interlineado "
            f"{paragraph.paragraph_format.line_spacing}, se esperaba {sample_config.line_spacing}"
        )

        # 2. Comprobación a nivel de OOXML (w:spacing lineRule="auto" y w:line)
        pPr = paragraph._element.find(qn("w:pPr"))
        assert pPr is not None, f"Párrafo {i} carece del elemento w:pPr"
        spacing = pPr.find(qn("w:spacing"))
        assert spacing is not None, f"Párrafo {i} carece del elemento w:spacing"

        assert spacing.get(qn("w:lineRule")) == "auto", (
            f"Párrafo {i}: w:lineRule esperado 'auto', obtenido '{spacing.get(qn('w:lineRule'))}'"
        )

        expected_line = str(round(sample_config.line_spacing * 240))
        assert spacing.get(qn("w:line")) == expected_line, (
            f"Párrafo {i}: w:line esperado '{expected_line}', obtenido '{spacing.get(qn('w:line'))}'"
        )


# ---------------------------------------------------------------------------
# Requisito 6: Test de preservación de tablas, imágenes y listas
# ---------------------------------------------------------------------------
def test_preservation_tables_images_and_lists(valid_docx_path, sample_config):
    """
    Verifica que el formateo conserve la estructura original:
    el conteo de tablas, imágenes (a:blip) y párrafos de lista debe ser idéntico
    antes y después de aplicar la plantilla.
    """
    doc_before = Document(str(valid_docx_path))
    tables_before = len(doc_before.tables)
    images_before = len(doc_before.element.xpath(".//a:blip"))
    lists_before = count_list_paragraphs(doc_before)

    doc_after = Document(str(valid_docx_path))
    apply_template(doc_after, sample_config)

    tables_after = len(doc_after.tables)
    images_after = len(doc_after.element.xpath(".//a:blip"))
    lists_after = count_list_paragraphs(doc_after)

    assert tables_after == tables_before, (
        f"Discrepancia en tablas para '{valid_docx_path.name}': antes={tables_before}, después={tables_after}"
    )
    assert images_after == images_before, (
        f"Discrepancia en imágenes (a:blip) para '{valid_docx_path.name}': antes={images_before}, después={images_after}"
    )
    assert lists_after == lists_before, (
        f"Discrepancia en listas para '{valid_docx_path.name}': antes={lists_before}, después={lists_after}"
    )


# ---------------------------------------------------------------------------
# Requisito 7: Test de integridad tras formateo
# ---------------------------------------------------------------------------
def test_integrity_save_and_reopen(loaded_doc, sample_config):
    """
    Verifica que tras aplicar la plantilla, el documento se guarde en un flujo
    de bytes y se pueda reabrir con Document() sin errores de corrupción.
    """
    apply_template(loaded_doc, sample_config)

    buffer = io.BytesIO()
    loaded_doc.save(buffer)
    buffer.seek(0)

    reopened_doc = Document(buffer)
    assert reopened_doc is not None, "El documento no pudo ser instanciado tras guardarse."
    assert len(reopened_doc.sections) == len(loaded_doc.sections), (
        "El número de secciones difiere tras reabrir el documento guardado."
    )

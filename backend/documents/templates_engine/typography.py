"""
Módulo de tipografía para templates_engine (S2T4).
Aplica fuente, tamaños y normaliza atributos de tema OOXML en estilos y runs.
"""

from docx.shared import Pt
from docx.oxml.ns import qn


def apply_typography(doc, font_name, font_size, heading_sizes=None):
    """
    Aplica tipografía a los estilos y runs del documento DOCX.

    Args:
        doc: Objeto docx.Document.
        font_name (str): Nombre de la fuente configurada (e.g. 'Times New Roman').
        font_size (int|float): Tamaño base de la fuente en puntos.
        heading_sizes (dict|None): Tamaños opcionales para títulos h1, h2, h3.
    """
    if heading_sizes is None:
        heading_sizes = {}

    target_styles = {
        "Normal": font_size,
        "Heading 1": heading_sizes.get("h1", round(font_size * 1.5 * 2) / 2),
        "Heading 2": heading_sizes.get("h2", round(font_size * 1.25 * 2) / 2),
        "Heading 3": heading_sizes.get("h3", round(font_size * 1.1 * 2) / 2),
    }

    # 1. Configurar estilos y remover atributos de temas
    for style_name, size in target_styles.items():
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = font_name
            if size:
                style.font.size = Pt(size)
            rPr = style.element.find(qn("w:rPr"))
            if rPr is not None:
                rFonts = rPr.find(qn("w:rFonts"))
                if rFonts is not None:
                    for attr in list(rFonts.attrib.keys()):
                        if "theme" in attr.lower():
                            del rFonts.attrib[attr]
                    rFonts.set(qn("w:ascii"), font_name)
                    rFonts.set(qn("w:hAnsi"), font_name)

    # 2. Configurar runs en párrafos
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = font_name
            if run._r.rPr is not None:
                rFonts = run._r.rPr.find(qn("w:rFonts"))
                if rFonts is not None:
                    for attr in list(rFonts.attrib.keys()):
                        if "theme" in attr.lower():
                            del rFonts.attrib[attr]
                    rFonts.set(qn("w:ascii"), font_name)
                    rFonts.set(qn("w:hAnsi"), font_name)

    # 3. Configurar runs en tablas
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = font_name
                        if run._r.rPr is not None:
                            rFonts = run._r.rPr.find(qn("w:rFonts"))
                            if rFonts is not None:
                                for attr in list(rFonts.attrib.keys()):
                                    if "theme" in attr.lower():
                                        del rFonts.attrib[attr]
                                rFonts.set(qn("w:ascii"), font_name)
                                rFonts.set(qn("w:hAnsi"), font_name)

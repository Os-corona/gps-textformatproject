from docx import Document
from docx.shared import Pt, Inches


def apply_runs(paragraph, runs):
    for run_data in runs:
        run = paragraph.add_run(run_data["text"])
        if run_data.get("bold"):
            run.bold = True
        if run_data.get("italic"):
            run.italic = True
        if run_data.get("font"):
            run.font.name = run_data["font"]
        if run_data.get("size"):
            run.font.size = Pt(run_data["size"])


def add_image_to_paragraph(paragraph, image_path, width_inches=None):
    run = paragraph.add_run()
    if width_inches:
        run.add_picture(image_path, width=Inches(width_inches))
    else:
        run.add_picture(image_path)


def apply_list_style(doc, paragraph, list_info):
    if not list_info or not list_info.get("is_list"):
        return
    # Mejora futura: distinguir viñeta/numerada con numbering.xml
    style_name = "List Number" if list_info.get("is_numbered") else "List Bullet"
    try:
        paragraph.style = doc.styles[style_name]
    except KeyError:
        pass


def add_table_to_doc(doc, table_data):
    rows = table_data["rows"]
    if not rows:
        return None
    num_rows = len(rows)
    num_cols = max(sum(c.get("h_span", 1) for c in row) for row in rows)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = "Table Grid"

    for i, row in enumerate(rows):
        col = 0  # columna de cuadrícula real
        for cell_data in row:
            span = cell_data.get("h_span", 1)
            v_merge = cell_data.get("v_merge")
            cell = table.cell(i, col)

            if v_merge == "continue":
                # se une con la celda de arriba (el texto queda en la ancla)
                table.cell(i - 1, col).merge(cell)
            else:
                if span > 1:
                    cell = cell.merge(table.cell(i, col + span - 1))
                cell.text = cell_data.get("text", "")

            col += span
    return table


def add_tables_to_doc(doc, tables_data):
    for table_data in tables_data:
        add_table_to_doc(doc, table_data)
    return doc


def rebuild_from_structure(blocks):
    doc = Document()

    for block in blocks:
        btype = block.get("type")

        if btype == "paragraph":
            paragraph = doc.add_paragraph()

            try:
                paragraph.style = doc.styles[block["style_name"]]
            except KeyError:
                pass

            apply_list_style(doc, paragraph, block.get("list_info"))
            apply_runs(paragraph, block.get("runs", []))

            images = block.get("has_inline_image")
            if images:
                for path in images:
                    add_image_to_paragraph(paragraph, path, width_inches=4)

        elif btype == "table":
            add_table_to_doc(doc, block)  # block debe traer "rows"

    return doc
from docx import Document

def compare_documents(path_a, path_b):
    doc_a = Document(path_a)
    doc_b = Document(path_b)

    report = {
        "text_diff": [],
        "style_diff": [],
        "structure_diff": [],
    }

    paras_a = [p for p in doc_a.paragraphs if p.text.strip()]
    paras_b = [p for p in doc_b.paragraphs if p.text.strip()]

    if len(paras_a) != len(paras_b):
        report["structure_diff"].append(
            f"Número de párrafos distinto: {len(paras_a)} vs {len(paras_b)}"
        )

    for i, (pa, pb) in enumerate(zip(paras_a, paras_b)):
        if pa.text != pb.text:
            report["text_diff"].append({
                "index": i, "original": pa.text, "rebuilt": pb.text
            })
        if pa.style.name != pb.style.name:
            report["style_diff"].append({
                "index": i, "original": pa.style.name, "rebuilt": pb.style.name
            })

    return report

def compare_tables(path_a, path_b):
    doc_a = Document(path_a)
    doc_b = Document(path_b)
    diffs = []

    if len(doc_a.tables) != len(doc_b.tables):
        diffs.append(f"Número de tablas distinto: {len(doc_a.tables)} vs {len(doc_b.tables)}")
        return diffs

    for t, (table_a, table_b) in enumerate(zip(doc_a.tables, doc_b.tables)):
        for i, (row_a, row_b) in enumerate(zip(table_a.rows, table_b.rows)):
            for j, (cell_a, cell_b) in enumerate(zip(row_a.cells, row_b.cells)):
                if cell_a.text != cell_b.text:
                    diffs.append(f"Tabla {t}, celda ({i},{j}): '{cell_a.text}' vs '{cell_b.text}'")

    return diffs
